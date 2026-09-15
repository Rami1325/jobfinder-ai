"""Extract raw text from an uploaded resume (DOCX or PDF).

Also handles LinkedIn's profile export ("Profile.pdf" via More → Save to PDF):
the export is two-column (narrow sidebar with Contact/Top Skills/Languages,
wide main column with Summary/Experience/Education), and naive line-based PDF
extraction interleaves the columns mid-line — sidebar skills bleed into the
summary. When a LinkedIn export is detected, the page is re-extracted per
column (measured from a real export: sidebar ends ~x=185/612, main starts
~x=220, so the split ratio hits the gutter's center), and
`clean_linkedin_profile_text` strips the per-page footers and UI labels.
Both steps are no-ops for anything that isn't a LinkedIn export.
"""
from __future__ import annotations

import io
import re
import zipfile

from app.config import get_settings
from app.llm.limits import KB, InputTooLarge

# "Page 1 of 3" footer LinkedIn stamps on every exported page.
_LI_PAGE_FOOTER_RE = re.compile(r"^\s*Page \d+ of \d+\s*$", re.MULTILINE)
# The "(LinkedIn)" suffix the export appends after the profile URL.
_LI_URL_LABEL_RE = re.compile(r"\(LinkedIn\)")


def is_linkedin_profile_export(text: str) -> bool:
    """True for text extracted from LinkedIn's 'Save to PDF' profile export:
    it always carries the profile URL plus per-page 'Page N of N' footers."""
    return "linkedin.com/in/" in text and bool(_LI_PAGE_FOOTER_RE.search(text))


def clean_linkedin_profile_text(text: str) -> str:
    """Strip LinkedIn-export noise (page footers, URL label); returns other
    text unchanged so it is safe to run on every upload."""
    if not is_linkedin_profile_export(text):
        return text
    cleaned = _LI_PAGE_FOOTER_RE.sub("", text)
    cleaned = _LI_URL_LABEL_RE.sub("", cleaned)
    # Collapse the blank runs the removals leave behind.
    return re.sub(r"\n{3,}", "\n\n", cleaned).strip()


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".docx"):
        return _extract_docx(data)
    if name.endswith(".pdf"):
        text = _extract_pdf(data)
        if is_linkedin_profile_export(text):
            columns = _extract_pdf_linkedin_columns(data)
            # Guard against layout drift: the re-extraction must carry (almost)
            # all the same text, or we keep the naive interleaved version.
            if len(columns) >= 0.8 * len(text):
                text = columns
        return clean_linkedin_profile_text(text)
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Upload a .docx, .pdf, or .txt resume.")


def _assert_docx_expansion(data: bytes) -> None:
    """Refuse a .docx that expands past `max_docx_uncompressed_mb`.

    The DOCX twin of `_assert_page_count`, and it exists for the same reason:
    /resume/upload and /tools/scan feed a user's file straight into this parser,
    and a free account costs nothing to make, so an unbounded expansion is a
    cheap way to exhaust a serverless instance. `max_upload_mb` does not cover
    it — that
    caps COMPRESSED bytes, and a .docx is a zip whose parts `python-docx`
    expands into an lxml tree in full before a single line of text exists.
    Measured: 0.298 MB of zip is 102.0 MB of `document.xml` (343:1), so the
    10 MB compressed cap admits ~3.4 GB of XML from unremarkable content.

    Measured from `ZipInfo.file_size`, summed across the WHOLE archive:

    - It is central-directory METADATA, so nothing is decompressed to read it —
      the bomb is refused for the cost of parsing a few dozen bytes per member.
    - The whole archive, not `word/document.xml` alone, because `PackageReader`
      reads every part: a bomb parked in `styles.xml` or `word/media/` costs
      exactly the same and would otherwise walk straight past a targeted check.

    THE DECLARED SUM IS AN UPPER BOUND, NOT A CLAIM WE TRUST — that was checked
    rather than assumed, because "the central directory can lie" is the obvious
    evasion. `ZipExtFile._read1` clamps every member to `data[:self._left]` with
    `_left` seeded from `file_size`, and `_update_crc` then fails the member at
    EOF. So a zip that under-declares cannot amplify: it fails CLOSED. Verified
    by rewriting all three central-directory entries of that 102 MB bomb to
    `file_size: 1` — `zipfile` and `python-docx` then both raise
    `BadZipFile("Bad CRC-32")` having inflated nothing, which `_extract_docx`
    turns into a 400, not a 500. A second pass that
    re-read the archive to "verify" the sum would therefore add no guarantee and
    would inflate every legitimate upload twice.
    """
    settings = get_settings()
    cap_mb = settings.max_docx_uncompressed_mb
    if cap_mb <= 0:
        return
    # THE FLOOR IS DERIVED, NOT WRITTEN DOWN TWICE. Images are already
    # compressed, so an image-heavy .docx expands about 1:1 and `max_upload_mb`
    # is itself the largest expansion a legitimate CV can reach — measured at
    # 10.3 MB against the 10 MB upload cap. A cap below that would refuse a real
    # photo-heavy resume the upload cap had just admitted, which is the guard
    # firing on legitimate input. Deriving the floor keeps the two numbers
    # coupled: raising `max_upload_mb` cannot silently turn this into that guard.
    cap_mb = max(cap_mb, 3 * settings.max_upload_mb)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        expanded = sum(i.file_size for i in zf.infolist())
    if expanded > cap_mb * KB * KB:
        # `InputTooLarge`, not the `ValueError` `_assert_page_count` raises:
        # this IS "a user-owned document exceeded its cap, before any model call
        # was made", so it takes the app-level 413 with the structured KB detail
        # `apiErrorMessage` already renders ("That CV is too large to process…
        # It was NOT saved"), rather than a bare 400 carrying an English
        # sentence. 413 also keeps a limit the user can act on out of Sentry,
        # which captures only 5xx.
        raise InputTooLarge("resume", size_kb=-(-expanded // KB), cap_kb=cap_mb * KB)


def _extract_docx(data: bytes) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        _assert_docx_expansion(data)
        doc = Document(io.BytesIO(data))
    except (zipfile.BadZipFile, PackageNotFoundError, KeyError) as e:
        # Not a zip, a zip that is not a Word package, or a member whose CRC
        # does not check out — a .doc, a PDF or an .odt renamed .docx is the
        # common one, and it is user error. It takes the same ValueError -> 400
        # path as "Unsupported file type" above, instead of the 500 all three
        # produced before this. `KeyError` is in the tuple because python-docx
        # does not normalise it: a readable zip MISSING `[Content_Types].xml`
        # (a plain .zip renamed .docx) dies inside `ZipFile.read` rather than as
        # `PackageNotFoundError`, and the scope is one statement — a KeyError
        # from building a Document out of user bytes is a malformed package,
        # never our bug. `InputTooLarge` is not in the tuple on purpose: it must
        # reach its own 413 handler.
        raise ValueError("That .docx could not be opened — re-save it from Word and try again.") from e
    lines: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            lines.append(para.text)
    # Capture text in tables too (some resumes use them for layout).
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _assert_page_count(pdf) -> None:  # noqa: ANN001 - pdfplumber.PDF
    """Refuse absurd page counts before extracting.

    `pdf.pages` is lazy, so counting is cheap while extraction is not — this is
    the cheapest place to stop a crafted PDF from burning an instance's CPU on
    either upload route (/resume/upload, /tools/scan). The ceiling is
    deliberately far above
    any real resume (a master CV runs ~30 rendered pages at the extreme).
    """
    limit = get_settings().max_pdf_pages
    if len(pdf.pages) > limit:
        raise ValueError(f"That PDF has too many pages ({len(pdf.pages)}); {limit} is the limit.")


def _extract_pdf(data: bytes) -> str:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        _assert_page_count(pdf)
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
    return "\n".join(parts)


# LinkedIn export pages are 612pt wide; the sidebar ends ~x=185 and the main
# column starts ~x=220, so 0.34 (~x=208) lands inside the gutter.
_LI_COLUMN_SPLIT_RATIO = 0.34


def _extract_pdf_linkedin_columns(data: bytes) -> str:
    """Re-extract a LinkedIn profile export column by column: all sidebar text
    (Contact, Top Skills, Languages, Certifications) first, then the main
    column (name, Summary, Experience, Education) — so neither stream bleeds
    into the other. Pages without sidebar text (page 2+) contribute only to
    the main stream."""
    import pdfplumber

    sidebar: list[str] = []
    main: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            split_x = page.width * _LI_COLUMN_SPLIT_RATIO
            left = page.crop((0, 0, split_x, page.height)).extract_text() or ""
            right = page.crop((split_x, 0, page.width, page.height)).extract_text() or ""
            if left.strip():
                sidebar.append(left)
            if right.strip():
                main.append(right)
    return "\n".join(sidebar + main)
